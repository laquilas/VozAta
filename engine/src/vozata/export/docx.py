"""DOCX export with python-docx: styled headings and lists, the transcript as a table."""

from __future__ import annotations

import io

import docx

from vozata.domain import ExportDocument

# minutes heading levels sit one below the "Minutes" section heading
_HEADING_LEVEL = {"h1": 2, "h2": 3, "h3": 4}


class DocxExporter:
    media_type = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
    extension = "docx"

    def export(self, document: ExportDocument) -> bytes:
        word = docx.Document()
        word.add_heading(document.title, level=0)
        for line in document.details:
            word.add_paragraph(line)
        word.add_heading(document.minutes_title, level=1)
        if not document.minutes_blocks:
            word.add_paragraph().add_run(document.no_minutes).italic = True
        for block in document.minutes_blocks:
            if block.kind in _HEADING_LEVEL:
                word.add_heading(block.text, level=_HEADING_LEVEL[block.kind])
            elif block.kind == "bullet":
                word.add_paragraph(block.text, style="List Bullet")
            elif block.kind == "numbered":
                word.add_paragraph(block.text, style="List Paragraph")
            else:
                word.add_paragraph(block.text)
        word.add_heading(document.transcript_title, level=1)
        if document.rows:
            table = word.add_table(rows=1, cols=3)
            table.style = "Table Grid"
            for cell, text in zip(table.rows[0].cells, document.columns, strict=True):
                cell.text = text
                cell.paragraphs[0].runs[0].bold = True
            for row in document.rows:
                cells = table.add_row().cells
                cells[0].text, cells[1].text, cells[2].text = row.time, row.speaker, row.text
        buffer = io.BytesIO()
        word.save(buffer)
        return buffer.getvalue()
