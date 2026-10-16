from __future__ import annotations

import dataclasses
import io

import docx

from tests.export.test_markdown import document
from vozata.export.docx import DocxExporter
from vozata.ports import DocumentExporter

MINUTES = """# Ata
## Decisões
A equipe decidiu adiar a versão beta.
- Lançamento na terça-feira
1. Ana atualiza as notas
### Próximos passos
Revisão na segunda-feira."""


def reopen(data: bytes) -> docx.document.Document:
    return docx.Document(io.BytesIO(data))


def test_satisfies_the_port() -> None:
    exporter = DocxExporter()

    assert isinstance(exporter, DocumentExporter)
    assert exporter.extension == "docx"
    assert exporter.media_type.startswith("application/vnd.openxmlformats-officedocument")


def test_minutes_become_styled_headings_lists_and_paragraphs() -> None:
    word = reopen(DocxExporter().export(document(MINUTES)))

    paragraphs = [(p.style.name, p.text) for p in word.paragraphs if p.text]
    assert paragraphs[:3] == [
        ("Title", "Reunião de 2026-10-16 09:00"),
        ("Normal", "Duração: 01:02:05"),
        ("Normal", "Idioma: Português"),
    ]
    assert ("Heading 1", "Ata") in paragraphs
    assert ("Heading 2", "Ata") in paragraphs  # the minutes' own "# Ata", one level down
    assert ("Heading 3", "Decisões") in paragraphs
    assert ("Normal", "A equipe decidiu adiar a versão beta.") in paragraphs
    assert ("List Bullet", "Lançamento na terça-feira") in paragraphs
    assert ("List Paragraph", "1. Ana atualiza as notas") in paragraphs
    assert ("Heading 4", "Próximos passos") in paragraphs
    assert ("Heading 1", "Transcrição") in paragraphs


def test_the_transcript_is_a_table_with_a_header_row() -> None:
    word = reopen(DocxExporter().export(document()))

    [table] = word.tables
    cells = [[cell.text for cell in row.cells] for row in table.rows]
    assert cells == [
        ["Hora", "Falante", "Texto"],
        ["00:00:02", "Eu", "Bom dia, ação!"],
        ["00:01:05", "Outros", "Oi."],
    ]


def test_missing_minutes_and_an_empty_transcript() -> None:
    empty = dataclasses.replace(document(minutes=None), rows=[], details=[])

    word = reopen(DocxExporter().export(empty))

    texts = [p.text for p in word.paragraphs]
    assert "Nenhuma ata foi gerada para esta sessão." in texts
    assert word.tables == []
