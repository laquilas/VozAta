from __future__ import annotations

import io

import pytest

from tests.api.conftest import Api
from tests.api.test_sessions import SESSION, _recorded


def test_markdown_export_downloads_minutes_and_transcript(api: Api) -> None:
    _recorded(api, minutes=True)

    response = api.client.get(f"/sessions/{SESSION}/export", params={"format": "md"})

    assert response.status_code == 200
    assert response.headers["content-type"] == "text/markdown; charset=utf-8"
    assert response.headers["content-disposition"] == f'attachment; filename="{SESSION}-minutes.md"'
    text = response.text
    assert text.startswith("# Meeting of 2026-10-15 20:23\n")
    assert "Language: English" in text
    assert "# Minutes" in text
    assert text.index("Hello.") < text.index("Hi.")  # transcript in time order


def test_a_renamed_meeting_exports_under_its_title(api: Api) -> None:
    _recorded(api, minutes=True)
    api.client.patch(f"/sessions/{SESSION}", json={"title": "Budget review"})

    text = api.client.get(f"/sessions/{SESSION}/export", params={"format": "md"}).text

    assert text.startswith("# Budget review\n")
    assert "Date: 2026-10-15 20:23" in text


def test_the_docx_export_carries_the_title(api: Api) -> None:
    from docx import Document

    _recorded(api, minutes=True)
    api.client.patch(f"/sessions/{SESSION}", json={"title": "Budget review"})

    response = api.client.get(f"/sessions/{SESSION}/export", params={"format": "docx"})

    paragraphs = [paragraph.text for paragraph in Document(io.BytesIO(response.content)).paragraphs]
    assert paragraphs[0] == "Budget review"


@pytest.mark.parametrize(("export_format", "magic"), [("docx", b"PK"), ("pdf", b"%PDF")])
def test_binary_formats_download_with_their_extension(
    api: Api, export_format: str, magic: bytes
) -> None:
    _recorded(api, minutes=True)

    response = api.client.get(f"/sessions/{SESSION}/export", params={"format": export_format})

    assert response.status_code == 200
    assert response.content.startswith(magic)
    assert response.headers["content-disposition"].endswith(f'-minutes.{export_format}"')


def test_markdown_is_the_default_and_a_session_without_minutes_still_exports(api: Api) -> None:
    _recorded(api)

    response = api.client.get(f"/sessions/{SESSION}/export")

    assert response.status_code == 200
    assert "No minutes were generated for this session." in response.text


def test_an_unknown_format_is_a_400_listing_the_formats(api: Api) -> None:
    _recorded(api)

    response = api.client.get(f"/sessions/{SESSION}/export", params={"format": "odt"})

    assert response.status_code == 400
    error = response.json()["error"]
    assert error["code"] == "unsupported_format"
    assert "docx, md, pdf" in error["message"]


def test_an_unknown_session_is_a_404(api: Api) -> None:
    response = api.client.get("/sessions/20991231_000000/export", params={"format": "pdf"})

    assert response.status_code == 404
    assert response.json()["error"]["code"] == "session_not_found"


def test_the_docx_export_opens(api: Api) -> None:
    import docx

    _recorded(api, minutes=True)

    content = api.client.get(f"/sessions/{SESSION}/export", params={"format": "docx"}).content

    word = docx.Document(io.BytesIO(content))
    assert word.paragraphs[0].text == "Meeting of 2026-10-15 20:23"


def test_the_export_uses_the_names_of_diarized_speakers(api: Api) -> None:
    from tests.api.test_sessions import _diarized

    _recorded(api, minutes=True)
    _diarized(api, {"system:1": "Bruno", "system:2": None})

    text = api.client.get(f"/sessions/{SESSION}/export", params={"format": "md"}).text

    assert "Bruno" in text
    assert "labels.user(n=2)" in text
